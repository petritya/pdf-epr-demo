import io
import zipfile
from openpyxl import load_workbook
from pypdf import PdfWriter
from fastapi.testclient import TestClient
import pytest
from converter import collect_pdfs, split_pages, convert, PageData, Table
from webapp import app

def pdf(n=1):
    writer=PdfWriter()
    for _ in range(n): writer.add_blank_page(width=100,height=100)
    out=io.BytesIO();writer.write(out);return out.getvalue()

def sample(_):
    return PageData(tables=[Table(title='Tételek',columns=['Kód','Érték'],rows=[['00123','=HYPERLINK("x")']],column_types=['text','text'],decimal_separator=',')],fields=[],other_text='Megjegyzés',warnings=['Olvasási bizonytalanság'])

def test_xlsx_preserves_identifiers_and_disables_formulas():
    name,data,mime=convert(split_pages(collect_pdfs([('a.pdf',pdf())])),extractor=sample)
    wb=load_workbook(io.BytesIO(data))
    assert name=='01_a.xlsx'
    assert wb['Tételek']['C2'].value=='00123'
    assert wb['Tételek']['D2'].data_type=='s'
    assert wb['Szöveg']['C2'].value=='Megjegyzés'
    assert wb['Ellenőrzések']['C3'].value=='Olvasási bizonytalanság'

def test_multiple_documents_distinct_outputs():
    _,data,mime=convert(split_pages(collect_pdfs([('a.pdf',pdf()),('a.pdf',pdf())])),extractor=sample)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert z.namelist()==['01_a.xlsx','02_a.xlsx']

def test_zip_traversal_rejected():
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:z.writestr('../a.pdf',pdf())
    with pytest.raises(ValueError):collect_pdfs([('a.zip',out.getvalue())])

def test_page_limit():
    with pytest.raises(ValueError):split_pages([('a.pdf',pdf(11))])

def test_invalid_pdf():
    with pytest.raises(ValueError):collect_pdfs([('a.pdf',b'not PDF')])

def test_valid_zip():
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:z.writestr('folder/a.pdf',pdf())
    assert len(split_pages(collect_pdfs([('a.zip',out.getvalue())])))==1

def test_unconfigured_service_cannot_send_data(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    with TestClient(app) as c:
        assert c.get('/').status_code==200
        assert not c.get('/config').json()['ready']
        assert c.post('/jobs',files={'files':('a.pdf',pdf(),'application/pdf')},data={'consent':'yes'}).status_code==503

def test_full_flow_with_fake_extractor(monkeypatch):
    import webapp
    monkeypatch.setenv('OPENAI_API_KEY','test-not-real')
    monkeypatch.setenv('PRIVACY_URL','https://example.com/privacy')
    monkeypatch.setattr(webapp,'convert',lambda docs,progress:convert(docs,extractor=sample,progress=progress))
    with TestClient(app) as c:
        res=c.post('/jobs',files={'files':('a.pdf',pdf(),'application/pdf')},data={'consent':'yes'})
        assert res.status_code==202
        token=res.json()['id']
        import time
        for _ in range(100):
            s=c.get('/jobs/'+token).json()
            if s['state']=='done':break
            time.sleep(.01)
        assert s['state']=='done'
        assert c.get('/jobs/'+token+'/download').status_code==200
        assert 'data' not in s
        assert c.get('/jobs/unknown/download').status_code==404

def test_typed_numbers_and_metadata():
    from converter import DocumentField, make_workbook
    page=PageData(tables=[Table(title='Tételek',columns=['Cikkszám','Mennyiség','Nettó összeg','Súly'],column_types=['text','number','number','number'],decimal_separator=',',rows=[['00123','2,00','15 740,00','0,96']])],fields=[DocumentField(label='Nettó összesen (HUF)',value='59 000,00 HUF',kind='text',decimal_separator=',')],other_text='',warnings=[])
    wb=load_workbook(io.BytesIO(make_workbook('minta.pdf',[page])))
    assert wb.sheetnames[0]=='Tételek'
    assert wb['Tételek'].max_row==2
    assert [wb['Tételek'].cell(2,i).value for i in range(3,7)]==['00123',2,15740,.96]
    assert wb['Dokumentumadatok']['D2'].value==59000
    assert wb['Dokumentumadatok']['E2'].value=='59 000,00 HUF'

@pytest.mark.parametrize('raw,sep,expected', [('1.234,56',',',1234.56),('1,234.56','.',1234.56),('-2,48',',',-2.48),('00123',',',None),('1234567890123456',',',None),('1.23,4',',',None),('=1+1',',',None),('27%',',',None)])
def test_numeric_parsing(raw,sep,expected):
    from converter import numeric_value
    result=numeric_value(raw,sep)
    assert (result[0] if result else None)==expected

def test_visible_image_input(monkeypatch):
    import converter
    monkeypatch.setenv('OPENAI_API_KEY','test')
    captured={}
    class Reply:
        def raise_for_status(self): pass
        def json(self):
            return {'status':'completed','output':[{'content':[{'type':'output_text','text':converter.DocumentData(pages=[sample(None)]).model_dump_json()}]}]}
    def post(*args,**kwargs):
        captured.update(kwargs['json']);return Reply()
    monkeypatch.setattr(converter.httpx,'post',post)
    converter.extract_page(pdf())
    part=captured['input'][0]['content'][2]
    assert part['type']=='input_image'
    assert part['image_url'].startswith('data:image/png;base64,')
    assert not captured['store']

def test_overlapping_text_objects_remain_separate():
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from converter import native_text
    import json
    writer=PdfWriter();page=writer.add_blank_page(width=500,height=200)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject()
    stream.set_data(b'BT /F1 10 Tf 1 0 0 1 10 100 Tm (A very long description overlapping code) Tj ET BT /F1 10 Tf 1 0 0 1 150 100 Tm (001-AB) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    out=io.BytesIO();writer.write(out)
    spans=json.loads(native_text(out.getvalue()))
    assert [s[2] for s in spans]==['A very long description overlapping code','001-AB']
    assert spans[0][0]<spans[1][0]

def test_document_rejects_missing_output_page(monkeypatch):
    import converter
    monkeypatch.setenv('OPENAI_API_KEY','test')
    class Reply:
        def raise_for_status(self):pass
        def json(self):return {'status':'completed','output':[{'content':[{'type':'output_text','text':converter.DocumentData(pages=[sample(None)]).model_dump_json()}]}]}
    monkeypatch.setattr(converter.httpx,'post',lambda *a,**k:Reply())
    with pytest.raises(RuntimeError,match='Nem minden oldal'):
        converter.extract_document([pdf(),pdf()])
